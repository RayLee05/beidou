"""Render the maintained architecture Markdown to Word using workspace dependencies.

Markdown is the content source. <!-- pagebreak --> controls logical pages;
running this script no longer reinstates an obsolete hardcoded architecture.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / '总体架构方案'
TITLE = '北斗B1I基带处理器_总体架构方案'


def font(run, size=12, bold=False):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('`', '').replace('**', '').replace('<br>', '\n')


def paragraph(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(.85 if size == 12 else 0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(18 if size == 12 else 15)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if 'SHA256' in text else WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(plain(text)), size)
    return p


def table(doc, rows):
    n = len(rows[0]); widths = [4.3, 11.7] if n == 2 else [3.0, 5.0, 8.0]
    if len(widths) != n: widths = [16 / n] * n
    t = doc.add_table(rows=0, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Cm(w)
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: el.set(qn('w:' + k), v)
        b.append(el)
    t._tbl.tblPr.append(b)
    margins = OxmlElement('w:tblCellMar')
    for side, val in [('top', '80'), ('bottom', '80'), ('left', '100'), ('right', '100')]:
        el = OxmlElement('w:' + side); el.set(qn('w:w'), val); el.set(qn('w:type'), 'dxa'); margins.append(el)
    t._tbl.tblPr.append(margins)
    for i, vals in enumerate(rows):
        row = t.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if not i: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, val, w in zip(row.cells, vals, widths):
            cell.width = Cm(w); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(14)
            font(p.add_run(plain(val.strip())), 10.5, i == 0)
            if not i:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(sh)
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = Pt(4); p.paragraph_format.space_after = Pt(0)


def build_docx(path):
    doc = Document(path) if path.exists() else Document()
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'): doc.element.body.remove(child)
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.15)
    sec.header_distance = Cm(.8); sec.footer_distance = Cm(.9)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        s = doc.styles[name]; s.font.name = 'Times New Roman'; s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体'); s.paragraph_format.first_line_indent = Cm(0)
    # Existing Word templates may put a theme-colored rule on Title.
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Title'].font.underline = False
    for level in (1, 2):
        s = doc.styles[f'Heading {level}']; s.font.size = Pt(14 if level == 1 else 12); s.font.bold = True
        s.paragraph_format.keep_with_next = True; s.paragraph_format.space_before = Pt(6)
        s.paragraph_format.space_after = Pt(4); s.paragraph_format.line_spacing = Pt(20)
    h = sec.header.paragraphs[0]; h.clear(); h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(h.add_run('北斗 B1I 基带接收机　总体架构'), 9)
    f = sec.footer.paragraphs[0]; f.clear(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run('第 '), 9); el = OxmlElement('w:fldSimple'); el.set(qn('w:instr'), 'PAGE'); f._p.append(el); font(f.add_run(' 页'), 9)
    lines = (OUT / (TITLE + '.md')).read_text(encoding='utf-8').splitlines(); i = 0; pagebreak = False
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line == '<!-- pagebreak -->': pagebreak = True; continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = Pt(30); p.paragraph_format.space_after = Pt(10); font(p.add_run(line[2:]), 22, True)
        elif line.startswith('##'):
            level = 2 if line.startswith('### ') else 1
            p = doc.add_paragraph(line.lstrip('# '), style=f'Heading {level}')
            p.paragraph_format.page_break_before = pagebreak; pagebreak = False
            for run in p.runs: font(run, 12 if level == 2 else 14, True)
        elif line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i += 1
            table(doc, [r.strip('|').split('|') for r in block if not re.fullmatch(r'[| :\-]+', r)])
        elif line.startswith('!['):
            cap, image = re.match(r'!\[(.*?)\]\((.*?)\)', line).groups()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(OUT / image), width=Cm(16))
            p = paragraph(doc, cap, 9.5); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('```'):
            code = []
            while i < len(lines) and not lines[i].startswith('```'): code.append(lines[i]); i += 1
            i += 1; p = paragraph(doc, '\n'.join(code), 9); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs: run.font.name = 'Consolas'
        else: paragraph(doc, line, 9.5 if line.startswith('版本：') else 12)
    doc.core_properties.title = '北斗 B1I 基带处理器总体架构方案'; doc.core_properties.author = '项目组'
    doc.core_properties.comments = '2026-10-13 修订；实现基线 2556e76；保留历史验证范围'
    doc.save(path); print(path)


if __name__ == '__main__': build_docx(OUT / (TITLE + '.docx'))
